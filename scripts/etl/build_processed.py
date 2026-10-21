# /// script
# requires-python = ">=3.12"
# dependencies = ["openpyxl>=3.1", "pyproj>=3.6"]
# ///
"""Converte os Excel recebidos da Vitrus em ficheiros limpos em data/processed/.

Uso: `just etl` (ou `uv run scripts/etl/build_processed.py`).
Os ficheiros em data/raw/ nunca se editam à mão; tudo o que é derivado sai daqui.

Os textos originais (em português) ficam como vieram; ao lado de cada categoria há uma chave estável
(`*_key`) em inglês, que a API expõe como enum e o web traduz.
"""

from __future__ import annotations

import csv
import hashlib
import json
import re
import unicodedata
from collections import defaultdict
from pathlib import Path

import openpyxl

import containers  # mesmo diretório (scripts/etl); `uv run` põe-no no sys.path
import textile_base as base  # base do têxtil gerado (Vitrus + APA), um só sítio no ETL

ROOT = Path(__file__).resolve().parents[2]
RAW = ROOT / "data" / "raw"
OUT = ROOT / "data" / "processed"

POPULATION = base.POPULATION
MONTHS = base.MONTHS

TYPOLOGY = {
    "Orgânicos": "organic",
    "Indiferenciados": "undifferentiated",
    "Recicláveis": "recyclable",
    "Volumosos (Monstros)": "bulky",
    "PAYT": "payt",
    "Serviços de apoio": "support",
    "A confirmar": "to_confirm",
    "Sem afetação": "unassigned",
}
COLLECTION = {"PAP": "door_to_door", "Moloks": "moloks", "Contentores": "containers", "": "unspecified"}
STATUS = {"OK": "ok", "Sem viatura atribuída": "no_vehicle", "Só ao fim de semana": "weekend_only", "—": "unassigned"}
DAY_GROUP = {"Semana (2ª a 6ª)": "weekday", "Sábado": "saturday", "Domingo": "sunday"}
SHIFT = {"Manhã": "morning", "Tarde": "afternoon", "Noite": "night"}
ALERT = {
    "Circuito sem viatura": "circuit_without_vehicle",
    "Código de viatura": "vehicle_code",
    "Circuito desconhecido": "unknown_circuit",
    "Circuito incompleto": "incomplete_circuit",
    "Viatura em duplicado": "duplicate_vehicle",
    "Possível erro de matrícula": "possible_plate_error",
    "Viatura multi-tipologia": "multi_typology_vehicle",
    "Viaturas sem circuito": "vehicle_without_circuit",
    "Circuitos não documentados": "undocumented_circuits",
    "Datas diferentes": "different_dates",
    "Escrita normalizada": "normalized_spelling",
}
WASTE = {
    "Biorresíduos": "biowaste",
    "Indiferenciados": "undifferentiated",
    "Papel": "paper",
    "Plástico": "plastic",
    "Vidro": "glass",
    "Madeira": "wood",
    "Metais": "metals",
    "outros": "others",
    "Compósitos": "composites",
    "Têxteis": "textiles",
    "Têxteis Sanitários": "sanitary_textiles",
    "Volumosos": "bulky",
}
CIRCUIT_SHEETS = ["Orgânicos", "Indiferenciados", "Recicláveis", "Volumosos", "Outros serviços"]
# Dias: 1=segunda ... 6=sábado, 7=domingo (ISO).
DAY_TOKENS = {"2ª": 1, "3ª": 2, "4ª": 3, "5ª": 4, "6ª": 5, "SAB": 6, "DOM": 7}
SITE_WORDS = re.compile(r"horeca|cemit|contentores|escolas|supermercado|todos os|nascional", re.I)


def _clean(value: object) -> str:
    return re.sub(r"\s+", " ", str(value)).strip() if value is not None else ""


def _write_csv(name: str, rows: list[dict]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    with (OUT / name).open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=list(rows[0]), lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def _days_from_text(text: str) -> set[int]:
    """Lê "2ª, 4ª & 6ª", "Todos os dias" ou "Segunda-feria Sexta-feira". Vazio se não houver dados."""
    low = text.lower()
    if "todos os dias" in low:
        return set(range(1, 8))
    days = {d for token, d in DAY_TOKENS.items() if re.search(rf"(?<!\w){re.escape(token)}", text, re.I)}
    if "segunda" in low:
        days.add(1)
    if "sexta" in low:
        days.add(5)
    return days


def _days_from_shifts(groups: set[str]) -> set[int]:
    days: set[int] = set()
    if "weekday" in groups:
        days |= {1, 2, 3, 4, 5}
    if "saturday" in groups:
        days.add(6)
    if "sunday" in groups:
        days.add(7)
    return days


# --- Têxtil e resíduos ---------------------------------------------------------------------------------------


def textile_flows() -> dict:
    """Destinos do têxtil recolhido por entidade em 2025, em formato longo (uma linha por destino)."""
    ws = openpyxl.load_workbook(RAW / "dados_textil.xlsx", data_only=True)["têxtil"]
    rows = []
    for entity, total, reuse, recycle, reject, source, *_ in ws.iter_rows(min_row=4, values_only=True):
        if not entity or str(entity).startswith(("TOTAL", "%")):
            continue
        method = "weighed" if "kg exatos" in (source or "") else "percentage"
        for destination, tonnes in (("reuse", reuse), ("recycle", recycle), ("reject", reject)):
            rows.append(
                {
                    "entity": entity,
                    "year": 2025,
                    "destination": destination,
                    "tonnes": round(tonnes, 4),
                    "method": method,
                    "provenance": _clean(source),
                }
            )
    _write_csv("textile_flows.csv", rows)
    total = sum(r["tonnes"] for r in rows)
    # Têxtil gerado = o que ainda vai para o lixo (43 051 t RSU × 5,4%, Vitrus + APA) + o recolhido à parte.
    demand_t = base.generated_t(total)
    by_dest = {
        d: round(sum(r["tonnes"] for r in rows if r["destination"] == d), 3) for d in ("reuse", "recycle", "reject")
    }
    return {
        "total_t": round(total, 3),
        **{f"{d}_t": v for d, v in by_dest.items()},
        "kg_per_inhab_month": round(total * 1000 / POPULATION / MONTHS, 3),
        "residual_kg_per_inhab_month": round(base.KG_PER_INHAB_MONTH, 4),
        "estimated_demand_t": round(demand_t, 1),
        "capture_of_demand_pct": round(base.capture(total) * 100, 1),
        "gap_t": round(demand_t - total, 1),
        "demand_source": base.SOURCE_LABEL,
    }


def waste_composition() -> None:
    """Composição dos RSU: colunas L:M da folha RSU, Europa primeiro e Portugal a seguir."""
    ws = openpyxl.load_workbook(RAW / "dados_textil.xlsx", data_only=True)["RSU"]
    rows: list[dict] = []
    scope = "europe"
    for row in ws.iter_rows(min_row=3, values_only=True):
        label = _clean(row[11]).replace("Biorrsesíduos", "Biorresíduos").replace("Compositos", "Compósitos")
        value = row[12]
        if not label or not isinstance(value, (int, float)):
            continue
        if label == "Biorresíduos" and rows:  # a segunda tabela (Portugal) recomeça em Biorresíduos
            scope = "portugal"
        rows.append({"scope": scope, "category_key": WASTE[label], "category": label, "share_pct": value})
    _write_csv("waste_composition.csv", rows)


# --- Circuitos, turnos e viaturas ------------------------------------------------------------------------------


def _localities(places: object) -> list[str]:
    parts = [p.strip() for p in re.split(r"[\n,]| e ", str(places or "")) if p.strip()]
    return [p for p in parts if not SITE_WORDS.search(p)]


def circuits() -> dict:
    wb = openpyxl.load_workbook(RAW / "Circuitos_Viaturas_Guimaraes.xlsx", data_only=True)

    # Turnos planeados (Base_Dados): uma linha por circuito, viatura e turno.
    shifts = []
    for row in list(wb["Base_Dados"].iter_rows(values_only=True))[1:]:
        day, shift, circuit, typology, plate, crew, kind, places, written, sheet, note = row[:11]
        if not circuit:
            continue
        shifts.append(
            {
                "day_group": DAY_GROUP[day],
                "shift": SHIFT[shift],
                "circuit": circuit,
                "typology_key": TYPOLOGY[typology],
                "plate": plate,
                "crew": int(crew or 0),
                "collection_key": COLLECTION[kind or ""],
                "localities": "; ".join(_localities(places)),
                "note": _clean(note),
            }
        )
    _write_csv("shifts.csv", shifts)

    groups_by_circuit: dict[str, set[str]] = defaultdict(set)
    for s in shifts:
        groups_by_circuit[s["circuit"]].add(s["day_group"])

    # Circuitos (folhas por tipologia) com dias de recolha e frequência semanal.
    circuit_rows = []
    for sheet in CIRCUIT_SHEETS:
        for row in wb[sheet].iter_rows(min_row=5, values_only=True):
            typology, circuit, kind, places, days, *_rest = row[:6]
            status, note = row[13], row[14]
            if not circuit or circuit == "SEM CIRCUITO":
                continue
            day_set = _days_from_text(_clean(days)) or _days_from_shifts(groups_by_circuit.get(circuit, set()))
            circuit_rows.append(
                {
                    "typology_key": TYPOLOGY[typology],
                    "circuit": circuit,
                    "collection_key": COLLECTION[kind or ""],
                    "localities": "; ".join(_localities(places)),
                    "days_text": _clean(days),
                    "days": " ".join(str(d) for d in sorted(day_set)),
                    "days_per_week": len(day_set),
                    "status_key": STATUS[status or "—"],
                    "note": _clean(note),
                }
            )
    _write_csv("circuits.csv", circuit_rows)

    # Cobertura por localidade: que circuitos passam, em que dias, com que tipo de recolha.
    by_locality: dict[str, dict] = {}
    for c in circuit_rows:
        for name in [x for x in c["localities"].split("; ") if x]:
            entry = by_locality.setdefault(name, {"circuits": [], "typologies": set(), "kinds": set(), "days": set()})
            entry["circuits"].append(c["circuit"])
            entry["typologies"].add(c["typology_key"])
            if c["collection_key"] != "unspecified":
                entry["kinds"].add(c["collection_key"])
            entry["days"] |= {int(d) for d in c["days"].split()}
    locality_rows = [
        {
            "locality": name,
            "circuits": "; ".join(sorted(set(e["circuits"]))),
            "typologies": "; ".join(sorted(e["typologies"])),
            "collection_kinds": "; ".join(sorted(e["kinds"])),
            "days": " ".join(str(d) for d in sorted(e["days"])),
            "days_per_week": len(e["days"]),
        }
        for name, e in sorted(by_locality.items())
    ]
    _write_csv("localities.csv", locality_rows)

    # Viaturas: a lista completa da folha Viaturas (inclui as que nunca têm circuito).
    vehicle_rows = [
        {"plate": row[0], "note": _clean(row[8])}
        for row in wb["Viaturas"].iter_rows(min_row=5, values_only=True)
        if row[0]
    ]
    _write_csv("vehicles.csv", vehicle_rows)

    # Capacidade da frota por tipologia ("Total de Viaturas" e "Total de RH" da folha Resumo, dados do Miguel).
    capacity_rows = [
        {"typology": TYPOLOGY[row[0]], "vehicles": int(row[10]), "crew": int(row[11])}
        for row in wb["Resumo"].iter_rows(min_row=5, values_only=True)
        if row[0] and row[0] != "TOTAL" and isinstance(row[10], (int, float)) and isinstance(row[11], (int, float))
    ]
    _write_csv("fleet_capacity.csv", capacity_rows)

    alerts = []
    for row in list(wb["Alertas"].iter_rows(min_row=4, values_only=True)):
        if row[0]:
            alerts.append(
                {
                    "id": row[0],
                    "type_key": ALERT[row[1]],
                    "detail": _clean(row[2]),
                    "suggestion": _clean(row[3]),
                }
            )
    _write_csv("data_alerts.csv", alerts)

    plates = {s["plate"] for s in shifts}
    return {
        "circuits": len(circuit_rows),
        "circuits_without_vehicle": sum(1 for c in circuit_rows if c["status_key"] == "no_vehicle"),
        "vehicles": len(plates),
        "vehicles_never_assigned": sorted(
            {v["plate"] for v in vehicle_rows} - {s["plate"] for s in shifts if s["circuit"] != "SEM CIRCUITO"}
        ),
        "shifts": len(shifts),
        "localities": len(locality_rows),
        "alerts": len(alerts),
    }


def geography() -> dict:
    """Dados públicos descarregados por `fetch_public_data.py`: população por freguesia e coordenadas."""
    public = RAW / "public"
    if not (public / "parishes.geojson").exists():
        return {"parishes": 0, "localities_located": 0}
    collection = json.loads((public / "parishes.geojson").read_text(encoding="utf-8"))
    tipau_file = public / "tipau.json"
    tipau = json.loads(tipau_file.read_text(encoding="utf-8")) if tipau_file.exists() else {}

    def norm(text: str) -> str:
        plain = "".join(c for c in unicodedata.normalize("NFD", text.lower()) if not unicodedata.combining(c))
        return re.sub(r"[^a-z0-9]+", " ", plain).strip()

    parish_rows = [
        {
            "dicofre": f["properties"]["dicofre"],
            "name": f["properties"]["name"],
            "population": f["properties"]["population"],
            "population_year": f["properties"]["year"],
            "tipau": tipau.get(norm(f["properties"]["name"]), ""),
            "latitude": f["properties"]["lat"],
            "longitude": f["properties"]["lon"],
            "geometry": json.dumps(f["geometry"], separators=(",", ":")) if f["geometry"] else "",
        }
        for f in collection["features"]
    ]
    _write_csv("parishes.csv", sorted(parish_rows, key=lambda r: r["dicofre"]))
    located = json.loads((public / "localities_geocoded.json").read_text(encoding="utf-8"))
    locality_rows = [
        {
            "locality": x["locality"],
            "latitude": x["lat"] or "",
            "longitude": x["lon"] or "",
            "dicofre": x["dicofre"] or "",
        }
        for x in located
    ]
    _write_csv("localities_geo.csv", locality_rows)
    return {"parishes": len(parish_rows), "localities_located": sum(1 for r in locality_rows if r["latitude"] != "")}


def sources() -> None:
    """Manifesto de proveniência: ficheiro, origem, período e impressão digital (SHA-256) de cada Excel."""

    def sha(name: str) -> str:
        return hashlib.sha256((RAW / name).read_bytes()).hexdigest()

    manifest = [
        {
            "key": "vitrus_textile_2025",
            "title": "Têxteis recolhidos em 2025 e composição dos RSU",
            "provider": "Vitrus Ambiente",
            "file_name": "dados_textil.xlsx",
            "file_sha256": sha("dados_textil.xlsx"),
            "received_on": "2026-10-19",
            "period_start": "2025-01-01",
            "period_end": "2025-12-31",
            "notes": "Sarah Trading e Ultriplo com kg exatos por destino; Nova Etapa e Vitrus com percentagens de relatórios.",
        },
        {
            "key": "vitrus_circuits_2026_09",
            "title": "Circuitos, viaturas e turnos de recolha",
            "provider": "Vitrus Ambiente",
            "file_name": "Circuitos_Viaturas_Guimaraes.xlsx",
            "file_sha256": sha("Circuitos_Viaturas_Guimaraes.xlsx"),
            "received_on": "2026-10-19",
            "period_start": "2026-09-19",
            "period_end": "2026-09-25",
            "notes": "Semana de 21 a 25/09/2026, sábado 19/09 e domingo 20/09 (fim de semana anterior).",
        },
    ]
    # Valores comunicados por email (sem ficheiro): a impressão digital é a dos próprios valores.
    vitrus_apa = json.dumps(
        {"rsu_t": base.VITRUS_RSU_2025_T, "textile_share": base.APA_TEXTILE_SHARE_OF_RSU}, sort_keys=True
    )
    manifest.append(
        {
            "key": "vitrus_apa_rsu_2025",
            "title": "RSU recolhido em 2025 e percentagem de têxtil no RSU",
            "provider": "Vitrus Ambiente (APA)",
            "file_name": "email do mentor da Vitrus (21/10/2026)",
            "file_sha256": hashlib.sha256(vitrus_apa.encode()).hexdigest(),
            "received_on": "2026-10-21",
            "period_start": "2025-01-01",
            "period_end": "2025-12-31",
            "notes": (
                "43 051 t de RSU recolhidas pela Vitrus em 2025; 5,4% é têxtil (estatística da APA, indicada pela "
                "Vitrus). Base do têxtil que ainda vai para o lixo (2 324,8 t/ano, estimativa)."
            ),
        }
    )
    public = RAW / "public"
    if (public / "parishes.geojson").exists():
        digest = hashlib.sha256(
            (public / "parishes.geojson").read_bytes() + (public / "localities_geocoded.json").read_bytes()
        ).hexdigest()
        manifest += [
            {
                "key": "ine_census_2021",
                "title": "População residente por freguesia (Censos 2021, indicador 0011609)",
                "provider": "INE, I.P.",
                "file_name": "data/raw/public/parishes.geojson",
                "file_sha256": digest,
                "received_on": "2026-10-19",
                "period_start": "2021-01-01",
                "period_end": "2021-12-31",
                "notes": "API pública do INE (pindica.jsp), um pedido por freguesia.",
            },
            {
                "key": "osm_nominatim",
                "title": "Contornos das freguesias e coordenadas das localidades",
                "provider": "OpenStreetMap contributors (ODbL)",
                "file_name": "data/raw/public/localities_geocoded.json",
                "file_sha256": digest,
                "received_on": "2026-10-19",
                "period_start": None,
                "period_end": None,
                "notes": "Geocodificação Nominatim. Algumas localidades dos circuitos (ex.: 'Ruas Estreitas') não são lugares e ficam sem coordenadas.",
            },
        ]
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "sources.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")


def main() -> None:
    summary = {"textile": textile_flows(), "circuits": circuits(), "geography": geography()}
    waste_composition()
    sources()
    summary["containers"] = containers.run()  # sem ficheiro em data/raw/: awaiting_data, não escreve nada
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
